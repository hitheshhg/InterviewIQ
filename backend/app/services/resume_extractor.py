import io
import logging
from typing import Tuple
from fastapi import UploadFile, HTTPException, status
import pdfplumber
import docx

logger = logging.getLogger("interviewiq.extractor")

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5MB
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


async def extract_text_from_file(file: UploadFile) -> Tuple[str, str]:
    """
    Safely extract plain text from an uploaded PDF or DOCX resume.
    Validates file type, size, corruption, empty contents, and scanned image status.
    Returns: (extracted_text, file_extension_without_dot)
    """
    filename = file.filename or "resume.pdf"
    file_ext = "." + filename.split(".")[-1].lower() if "." in filename else ""

    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{file_ext}'. Please upload a PDF (.pdf) or Word document (.docx).",
        )

    # Read contents into memory safely
    content = await file.read()

    # 1. Empty file validation
    if len(content) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty (0 bytes). Please upload a valid resume.",
        )

    # 2. File size validation
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum allowed size of 5MB (Uploaded: {round(len(content) / (1024 * 1024), 2)}MB).",
        )

    raw_text = ""

    # 3. Corrupt / unreadable file validation
    try:
        if file_ext == ".pdf":
            raw_text = _extract_from_pdf(content)
        elif file_ext == ".docx":
            raw_text = _extract_from_docx(content)
    except Exception as e:
        logger.warning(f"[Extractor] Corrupt or unreadable file '{filename}': {e}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file appears to be corrupted or unreadable. Please check the file and try again.",
        )

    # 4. Scanned image PDF / no selectable text layer validation
    cleaned = raw_text.strip()
    if len(cleaned) < 50:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="This document appears to be a scanned image or empty page without selectable text. Please upload a text-based PDF or DOCX resume.",
        )

    return cleaned, file_ext.replace(".", "")


def _extract_from_pdf(content: bytes) -> str:
    """Extract text from PDF bytes using pdfplumber."""
    # Basic PDF magic header check
    if not content.startswith(b"%PDF"):
        raise ValueError("Invalid PDF header signature.")

    text_parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        if not pdf.pages:
            return ""
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_from_docx(content: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    doc = docx.Document(io.BytesIO(content))
    text_parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    text_parts.append(cell.text)
    return "\n".join(text_parts)
